"""Reproducible UCI ingestion. Raw bytes and source rows are never overwritten.

Definitions: positive-priced, positive-quantity, non-C invoice lines are sales.
Exact duplicate source rows are retained: there is no evidence they are errors.
"""
import argparse
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import shutil
import zipfile
import duckdb
import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "datasets" / "uci_online_retail_ii"
URL = "https://archive.ics.uci.edu/static/public/502/online%2Bretail%2Bii.zip"
PAGE = "https://archive.ics.uci.edu/dataset/502/online+retail+ii"
RAW_DDL = """CREATE TABLE raw_lines (
 line_id BIGINT, source_sheet VARCHAR, source_row INTEGER, invoice_no VARCHAR,
 stock_code VARCHAR, description VARCHAR, quantity BIGINT, invoice_ts TIMESTAMP,
 unit_price DECIMAL(20,6), customer_id VARCHAR, country VARCHAR)"""


def digest(path):
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def download():
    DATA.mkdir(parents=True, exist_ok=True)
    archive = DATA / "online_retail_ii.zip"
    manifest = DATA / "source_manifest.json"
    if manifest.exists():
        info = json.loads(manifest.read_text(encoding="utf-8"))
        if digest(archive) != info["zip_sha256"] or digest(DATA / "online_retail_II.xlsx") != info["xlsx_sha256"]:
            raise RuntimeError("EXISTING_SOURCE_HASH_MISMATCH")
        print("Verified cached official source; no download.", flush=True)
        return
    if archive.exists() or (DATA / "online_retail_II.xlsx").exists():
        raise RuntimeError("UNMANIFESTED_SOURCE_NO_OVERWRITE")
    partial = DATA / "download.partial"
    with httpx.stream("GET", URL, timeout=httpx.Timeout(120, connect=20), follow_redirects=True) as response:
        response.raise_for_status()
        if response.url.host != "archive.ics.uci.edu":
            raise RuntimeError("UNEXPECTED_DOWNLOAD_HOST")
        size = 0
        with partial.open("wb") as handle:
            for chunk in response.iter_bytes(1024 * 1024):
                size += len(chunk)
                if size > 150_000_000:
                    raise RuntimeError("DOWNLOAD_TOO_LARGE")
                handle.write(chunk)
    with zipfile.ZipFile(partial) as zipped:
        members = [item for item in zipped.infolist() if Path(item.filename).name == "online_retail_II.xlsx"]
        if len(members) != 1 or members[0].file_size > 150_000_000:
            raise RuntimeError("UNEXPECTED_ARCHIVE_CONTENTS")
        workbook = DATA / "online_retail_II.xlsx"
        with zipped.open(members[0]) as source, workbook.open("xb") as output:
            shutil.copyfileobj(source, output)
    partial.rename(archive)
    info = {"source_url": URL, "source_page": PAGE, "doi": "10.24432/C5CG6D",
            "citation": "Chen, D. (2012). Online Retail II. UCI Machine Learning Repository.",
            "license": "CC BY 4.0", "downloaded_at_utc": datetime.now(timezone.utc).isoformat(),
            "zip_bytes": archive.stat().st_size, "zip_sha256": digest(archive), "xlsx_sha256": digest(workbook),
            "publisher_stated_rows": 1067371, "checksum_note": "Locally computed provenance hashes, not publisher-provided checksums."}
    manifest.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Official UCI source downloaded: {size} bytes; SHA-256 recorded.", flush=True)


def identifier(value):
    if value is None:
        return None
    if isinstance(value, (float, int)) and int(value) == value:
        return str(int(value))
    return str(value).strip() or None


def normalize_row(values, row_id, sheet, source_row):
    invoice, stock, description, quantity, stamp, price, customer, country = values[:8]
    if stamp is not None and not isinstance(stamp, datetime):
        raise ValueError("UNEXPECTED_INVOICE_DATE_TYPE")
    if quantity is not None and int(quantity) != quantity:
        raise ValueError("NON_INTEGRAL_QUANTITY")
    price = str(Decimal(str(price))) if price is not None else None
    return [row_id, sheet, source_row, identifier(invoice), identifier(stock),
            str(description).strip() if description is not None else None,
            int(quantity) if quantity is not None else None,
            stamp.isoformat(sep=" ") if stamp else None, price, identifier(customer), identifier(country)]


def build_warehouse(db):
    db.execute("""CREATE TABLE fact_retail_line AS
        SELECT line_id, invoice_no, stock_code, customer_id,
          CAST(strftime(invoice_ts, '%Y%m%d') AS INTEGER) AS date_id,
          country, quantity, unit_price,
          CAST(quantity * unit_price AS DECIMAL(24,6)) AS signed_amount,
          CAST(CASE WHEN invoice_no IS NOT NULL AND upper(invoice_no) NOT LIKE 'C%'
                 AND quantity > 0 AND unit_price > 0 THEN quantity * unit_price ELSE 0 END AS DECIMAL(24,6)) AS sale_amount,
          CASE WHEN invoice_no IS NOT NULL AND upper(invoice_no) NOT LIKE 'C%'
                 AND quantity > 0 AND unit_price > 0 THEN quantity ELSE 0 END AS sale_units,
          CASE WHEN invoice_no IS NOT NULL AND upper(invoice_no) NOT LIKE 'C%'
                 AND quantity > 0 AND unit_price > 0 THEN invoice_no END AS sale_invoice,
          CASE WHEN invoice_no IS NOT NULL AND upper(invoice_no) NOT LIKE 'C%'
                 AND quantity > 0 AND unit_price > 0 THEN customer_id END AS sale_customer,
          CAST(CASE WHEN upper(invoice_no) LIKE 'C%' AND quantity < 0 AND unit_price > 0
                 THEN -quantity * unit_price ELSE 0 END AS DECIMAL(24,6)) AS cancel_amount,
          CASE WHEN upper(invoice_no) LIKE 'C%' AND quantity < 0 AND unit_price > 0 THEN invoice_no END AS cancel_invoice
        FROM raw_lines""")
    db.execute("""CREATE TABLE dim_date AS SELECT DISTINCT date_id,
        CAST(substr(CAST(date_id AS VARCHAR),1,4) AS INTEGER) AS year,
        CAST(substr(CAST(date_id AS VARCHAR),5,2) AS INTEGER) AS month,
        CAST(substr(CAST(date_id AS VARCHAR),7,2) AS INTEGER) AS day,
        substr(CAST(date_id AS VARCHAR),1,4) || '-' || substr(CAST(date_id AS VARCHAR),5,2) AS year_month
        FROM fact_retail_line WHERE date_id IS NOT NULL""")
    db.execute("CREATE TABLE dim_country AS SELECT DISTINCT country FROM raw_lines WHERE country IS NOT NULL")
    db.execute("""CREATE TABLE dim_product AS SELECT stock_code, description AS product_name FROM raw_lines
        WHERE stock_code IS NOT NULL
        QUALIFY ROW_NUMBER() OVER (PARTITION BY stock_code ORDER BY (description IS NULL OR description=''),invoice_ts DESC,line_id DESC)=1""")


def profile(db):
    query = """SELECT count(*) AS raw_rows, count(DISTINCT invoice_no) AS invoices,
      count(DISTINCT stock_code) AS stock_codes, count(DISTINCT country) AS countries,
      count(*) FILTER (WHERE customer_id IS NULL) AS missing_customers,
      count(*) FILTER (WHERE description IS NULL OR description='') AS missing_descriptions,
      count(*) FILTER (WHERE invoice_ts IS NULL) AS missing_dates,
      count(*) FILTER (WHERE upper(invoice_no) LIKE 'C%') AS cancellation_rows,
      count(*) FILTER (WHERE quantity < 0) AS negative_quantity_rows,
      count(*) FILTER (WHERE unit_price <= 0) AS nonpositive_price_rows,
      min(invoice_ts) AS min_date, max(invoice_ts) AS max_date FROM raw_lines"""
    result = dict(zip([c[0] for c in db.execute(query).description], db.fetchone()))
    result["duplicate_extra_rows"] = db.execute("""SELECT sum(n-1) FROM (
      SELECT count(*) n FROM raw_lines GROUP BY invoice_no,stock_code,description,quantity,invoice_ts,unit_price,customer_id,country HAVING count(*)>1)""").fetchone()[0] or 0
    result["curated_rows"] = db.execute("SELECT count(*) FROM fact_retail_line").fetchone()[0]
    joined = db.execute("""SELECT count(*) FROM fact_retail_line f
      LEFT JOIN dim_date d ON f.date_id=d.date_id LEFT JOIN dim_country c ON f.country=c.country
      LEFT JOIN dim_product p ON f.stock_code=p.stock_code""").fetchone()[0]
    assert result["curated_rows"] == result["raw_rows"] == joined
    result["join_row_count"] = joined
    result["duplicate_policy"] = "Preserve all source rows; flag identical rows, do not silently deduplicate."
    result["currency"] = "GBP"
    result["timezone"] = "Source does not specify timezone; preserve source wall-clock time without conversion."
    result["monthly"] = [dict(zip(["month", "rows", "missing_customers"], row)) for row in db.execute("SELECT strftime(invoice_ts,'%Y-%m'),count(*),count(*) FILTER (WHERE customer_id IS NULL) FROM raw_lines GROUP BY 1 ORDER BY 1").fetchall()]
    return result


def load_normalized(db, csv_path):
    # Python csv.writer uses doubled quotes, never a backslash escape.
    db.execute("""INSERT INTO raw_lines SELECT * FROM read_csv(?,header=false,auto_detect=false,
      delim=',',quote='"',escape='"',new_line='\\r\\n',
      columns={'line_id':'BIGINT','source_sheet':'VARCHAR','source_row':'INTEGER','invoice_no':'VARCHAR','stock_code':'VARCHAR','description':'VARCHAR','quantity':'BIGINT','invoice_ts':'TIMESTAMP','unit_price':'DECIMAL(20,6)','customer_id':'VARCHAR','country':'VARCHAR'},nullstr='')""", [str(csv_path)])


def prepare():
    manifest = json.loads((DATA / "source_manifest.json").read_text(encoding="utf-8"))
    if digest(DATA / "online_retail_II.xlsx") != manifest["xlsx_sha256"]:
        raise RuntimeError("SOURCE_HASH_MISMATCH")
    target = DATA / "reference.duckdb"
    if target.exists():
        raise RuntimeError("REFERENCE_EXISTS_NO_OVERWRITE")
    csv_path = DATA / "normalized_rows.csv"
    if csv_path.exists():
        raise RuntimeError("NORMALIZED_FILE_EXISTS_NO_OVERWRITE")
    workbook = load_workbook(DATA / "online_retail_II.xlsx", read_only=True, data_only=True)
    total, sheets = 0, {}
    with csv_path.open("x", newline="", encoding="utf-8") as output:
        writer = csv.writer(output)
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = tuple(next(rows))[:8]
            if header != ("Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country"):
                raise RuntimeError("SOURCE_SCHEMA_CHANGED")
            count = 0
            for source_row, values in enumerate(rows, 2):
                if all(value is None for value in values):
                    continue
                total += 1
                count += 1
                writer.writerow(normalize_row(values, total, sheet.title, source_row))
                if total % 100000 == 0:
                    print(f"Read {total:,} source rows", flush=True)
            sheets[sheet.title] = count
    workbook.close()
    assert total == manifest["publisher_stated_rows"], "SOURCE_ROW_COUNT_CHANGED"
    db = duckdb.connect(str(target))
    db.execute(RAW_DDL)
    load_normalized(db, csv_path)
    build_warehouse(db)
    quality = profile(db)
    quality.update({"sheets": sheets, "source_xlsx_sha256": manifest["xlsx_sha256"], "normalized_csv_sha256": digest(csv_path), "definitions_version": "uci_retail_ii_v1"})
    (DATA / "quality.json").write_text(json.dumps(quality, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    db.close()
    print(json.dumps({k: v for k, v in quality.items() if k != "monthly"}, ensure_ascii=True, default=str), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["download", "prepare"])
    args = parser.parse_args()
    {"download": download, "prepare": prepare}[args.action]()
